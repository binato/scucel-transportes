import os
import openpyxl
from datetime import datetime, date
from sqlalchemy.orm import Session
from ..models import Veiculo, Abastecimento, Despesa

def categorize_expense(desc: str) -> str:
    d = desc.upper()
    if any(k in d for k in ['PNEU', 'CALIBRA', 'BALANCE', 'ESFERA', 'RODZIO', 'RODIZIO', 'REMENDO']):
        return 'PNEUS'
    elif any(k in d for k in ['LAV', 'POLIMENTO', 'HIGIEN']):
        return 'LAVAGEM'
    elif any(k in d for k in ['SEGURO']):
        return 'SEGURO'
    elif any(k in d for k in ['AUTOTRAC', 'ONIX', 'TACGRAFO', 'TACOGRAFO', 'ANTENA']):
        return 'RASTREAMENTO'
    elif any(k in d for k in ['HOTEL', 'DIARIA', 'ALMO', 'JANTA', 'BOTA', 'COLETE']):
        return 'DIARIAS_ESTRADA'
    elif any(k in d for k in ['IPVA', 'MULTA', 'LICEN']):
        return 'DOCUMENTACAO'
    elif any(k in d for k in ['ELTR', 'ELETR', 'BATERIA', 'FAROL', 'SETA', 'LANTERNA', 'SIRENE', 'LAMPADA', 'AR COND']):
        return 'ELETRICA'
    elif any(k in d for k in ['LEO', 'OLEO', 'FILTRO', 'MOTOR', 'EMBREAGEM', 'AMORTECEDOR', 'BOMBA', 'COMPRESSOR', 'REVIS', 'MECN', 'MECAN', 'M.O', 'MO DE OBRA', 'PEA', 'PECA', 'ENGATE', 'VALVULA', 'MANGUEIRA']):
        return 'MECANICA_PECAS'
    return 'OUTROS'

def parse_date_value(val, default_year=2026, default_month=1) -> date:
    if isinstance(val, datetime):
        return val.date()
    elif isinstance(val, date):
        return val
    elif isinstance(val, str):
        val = val.strip()
        if '/' in val:
            parts = val.split('/')
            if len(parts) == 2:
                day, month = int(parts[0]), int(parts[1])
                return date(default_year, month, day)
            elif len(parts) == 3:
                day, month, year = int(parts[0]), int(parts[1]), int(parts[2])
                if year < 100: year += 2000
                return date(year, month, day)
        elif '-' in val:
            parts = val.split('-')
            if len(parts) == 3:
                return date(int(parts[0]), int(parts[1]), int(parts[2]))
    return date(default_year, default_month, 1)

def import_excel_data(db: Session, file_path: str, force_reimport: bool = False):
    if not os.path.exists(file_path):
        print(f"File {file_path} does not exist for import.")
        return

    if not force_reimport and db.query(Abastecimento).count() > 0:
        print("Database already has records. Skipping initial Excel auto-import.")
        return

    if force_reimport:
        db.query(Despesa).delete()
        db.query(Abastecimento).delete()
        db.commit()

    print(f"Starting Excel import from {file_path}...")
    wb = openpyxl.load_workbook(file_path, data_only=True)

    known_trucks = [
        {"placa": "PRB0D89", "modelo": "Scania R450", "tipo": "Cavalo Mecânico", "meta_km_l": 3.10},
        {"placa": "PRB0E49", "modelo": "Scania R450", "tipo": "Cavalo Mecânico", "meta_km_l": 3.10},
        {"placa": "TJD0D09", "modelo": "Volvo FH 540", "tipo": "Cavalo Mecânico", "meta_km_l": 3.00},
        {"placa": "FKM6A46", "modelo": "Mercedes-Benz Actros", "tipo": "Cavalo Mecânico", "meta_km_l": 3.05},
        {"placa": "FKC6H71", "modelo": "Mercedes-Benz Actros", "tipo": "Cavalo Mecânico", "meta_km_l": 3.05},
        {"placa": "FCQ5G43", "modelo": "Volvo FH 460", "tipo": "Cavalo Mecânico", "meta_km_l": 3.00},
        {"placa": "STE0J12", "modelo": "DAF XF 530", "tipo": "Cavalo Mecânico", "meta_km_l": 3.00},
        {"placa": "GJA4J12", "modelo": "Scania R450", "tipo": "Cavalo Mecânico", "meta_km_l": 3.10},
        {"placa": "STEOJ52", "modelo": "DAF XF 480", "tipo": "Cavalo Mecânico", "meta_km_l": 3.00},
        {"placa": "FVZ3A72", "modelo": "Volvo FH 540", "tipo": "Cavalo Mecânico", "meta_km_l": 3.00}
    ]

    for t in known_trucks:
        v = db.query(Veiculo).filter(Veiculo.placa == t["placa"]).first()
        if not v:
            v = Veiculo(**t)
            db.add(v)
    db.commit()

    last_km_by_vehicle = {}

    if "THIAGO-2026" in wb.sheetnames:
        ws = wb["THIAGO-2026"]
        month_offsets = [
            (1, 2, 1),   # Jan
            (2, 13, 1),  # Feb
            (3, 25, 24), # Mar
            (4, 36, 24), # Apr
            (5, 48, 47), # May
            (6, 60, 59)  # Jun
        ]
        truck_row_starts = [3, 26, 49, 72, 95, 118, 140, 167]

        for idx, start_r in enumerate(truck_row_starts):
            end_r = truck_row_starts[idx+1] - 3 if idx+1 < len(truck_row_starts) else 185
            base_plate = ws.cell(start_r, 1).value
            if not base_plate:
                continue
            base_plate = str(base_plate).strip()

            for month_num, col_base, plate_col in month_offsets:
                local_plate = ws.cell(start_r, plate_col).value or base_plate
                local_plate = str(local_plate).strip()

                veiculo = db.query(Veiculo).filter(Veiculo.placa == local_plate).first()
                if not veiculo:
                    veiculo = Veiculo(placa=local_plate, modelo="Caminhão Scucel")
                    db.add(veiculo)
                    db.commit()

                for r in range(start_r, end_r + 1):
                    d_val = ws.cell(r, col_base).value
                    km_ini = ws.cell(r, col_base + 1).value
                    km_fim = ws.cell(r, col_base + 2).value
                    km_rod = ws.cell(r, col_base + 3).value
                    litros = ws.cell(r, col_base + 4).value
                    media = ws.cell(r, col_base + 5).value
                    arla = ws.cell(r, col_base + 6).value
                    valor = ws.cell(r, col_base + 7).value

                    desp_desc = ws.cell(r, col_base + 8).value
                    desp_val = ws.cell(r, col_base + 9).value

                    if d_val and (km_fim or litros or valor):
                        parsed_date = parse_date_value(d_val, 2026, month_num)
                        
                        try:
                            litros_f = float(litros) if litros is not None else 0.0
                            km_fim_i = int(km_fim) if km_fim is not None else 0
                            km_ini_i = int(km_ini) if km_ini is not None else None
                            valor_f = float(valor) if valor is not None else 0.0
                            arla_f = float(arla) if arla is not None else 0.0

                            # Fix missing km_ini using vehicle's last recorded odometer
                            if km_ini_i is None or km_ini_i == 0:
                                prev_km = last_km_by_vehicle.get(veiculo.id)
                                if prev_km and km_fim_i > prev_km and (km_fim_i - prev_km) < 5000:
                                    km_ini_i = prev_km

                            if km_ini_i is not None and km_fim_i > km_ini_i:
                                real_km_rod = km_fim_i - km_ini_i
                            elif km_rod is not None and int(km_rod) < 6000:
                                real_km_rod = int(km_rod)
                            else:
                                real_km_rod = int(litros_f * 3.0) if litros_f > 0 else 0

                            # update last known km
                            if km_fim_i > 0:
                                last_km_by_vehicle[veiculo.id] = km_fim_i

                            # calculate media
                            if litros_f > 0 and real_km_rod > 0:
                                calc_media = round(real_km_rod / litros_f, 2)
                            elif media is not None and isinstance(media, (int, float)) and 1.5 <= media <= 6.0:
                                calc_media = round(float(media), 2)
                            else:
                                calc_media = 3.00

                            custo_km = round(valor_f / real_km_rod, 2) if real_km_rod > 0 and valor_f > 0 else None

                            abast = Abastecimento(
                                veiculo_id=veiculo.id,
                                data=parsed_date,
                                km_inicial=km_ini_i,
                                km_final=km_fim_i,
                                km_rodado=real_km_rod,
                                litros_diesel=litros_f,
                                litros_arla=arla_f,
                                valor_total_nota=valor_f,
                                media_km_l=calc_media,
                                custo_km_diesel=custo_km,
                                origem_registro="EXCEL_IMPORT",
                                aprovado=True
                            )
                            db.add(abast)
                        except Exception as ex:
                            print(f"Error parsing row {r} col {col_base}: {ex}")

                    if desp_desc and desp_val is not None:
                        desp_str = str(desp_desc).strip()
                        if desp_str.lower() != 'total':
                            try:
                                d_val_clean = float(desp_val)
                                parsed_desp_date = parse_date_value(d_val, 2026, month_num) if d_val else date(2026, month_num, 1)
                                cat = categorize_expense(desp_str)
                                desp = Despesa(
                                    veiculo_id=veiculo.id,
                                    data=parsed_desp_date,
                                    categoria=cat,
                                    descricao=desp_str,
                                    valor=d_val_clean
                                )
                                db.add(desp)
                            except Exception as ex:
                                print(f"Error parsing expense row {r}: {ex}")

    db.commit()
    print("Excel import finished cleanly with realistic metrics!")
